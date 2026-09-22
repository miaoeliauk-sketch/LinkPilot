# -*- coding: utf-8 -*-
"""
全链路体检：一条命令查清楚"现取地址"这条路到底断在哪一环。

用法：
    python3 diagnose.py 你的表格.xlsx /路径/cookies.txt

会依次检查：
    1. cookies.txt 能不能读、有没有登录凭证
    2. 表格里能不能取到作品 id
    3. 调抖音接口能不能换到一个**新地址**（这一步就是"现取"）
    4. 这个新地址能不能真的下载（带 Referer）
每一步失败都会直接说清楚是什么问题、该怎么办。
"""
import os
import sys
import time
import urllib.error
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

REFERER = "https://www.douyin.com/"
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/146.0.0.0 Safari/537.36")


def die(msg):
    print(f"\n❌ {msg}")
    sys.exit(1)


def step(n, title):
    print(f"\n[{n}/4] {title}")


def find_row(xlsx):
    """从表格里找第一条能用的（作品id, 表格里存的地址）。"""
    from openpyxl import load_workbook
    import douyin_api

    ws = load_workbook(xlsx, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "") for c in next(rows)]

    def idx(name):
        return header.index(name) if name in header else None

    i_id, i_link = idx("作品id"), idx("作品网址")
    i_media = idx("视频源网址")
    if i_id is None and i_link is None:
        die(f"表格里既没有「作品id」也没有「作品网址」列。实际列名：{header}")

    for row in rows:
        def cell(i):
            return row[i] if i is not None and i < len(row) else None
        aweme_id = (douyin_api.extract_aweme_id(cell(i_id))
                    or douyin_api.extract_aweme_id(cell(i_link)))
        if aweme_id:
            return aweme_id, cell(i_media)
    die("表格里一条作品 id 都取不到。")


def probe(url):
    req = urllib.request.Request(
        url, headers={"User-Agent": UA, "Referer": REFERER}, method="GET")
    try:
        with urllib.request.urlopen(req, timeout=20) as r:
            return r.status, (r.headers.get("Content-Type") or "")
    except urllib.error.HTTPError as e:
        return e.code, ""
    except Exception as e:  # noqa: BLE001
        return None, str(e)


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    xlsx, cookies = sys.argv[1], sys.argv[2]

    try:
        import douyin_api
    except ImportError as e:
        die(f"导入 douyin_api 失败：{e}\n   说明你运行的不是最新版文件夹，或者没装依赖。")

    for path, what in ((xlsx, "表格"), (cookies, "cookies 文件")):
        if not os.path.isfile(path):
            die(f"找不到{what}：{path}")

    # 1. cookies
    step(1, "检查 cookies.txt ...")
    try:
        header = douyin_api.cookie_header_from_file(cookies)
    except douyin_api.DouyinApiError as e:
        die(f"{e}\n   → 用浏览器扩展（Get cookies.txt LOCALLY）登录抖音后重新导出。")
    names = [p.split("=")[0] for p in header.split("; ")]
    print(f"      读到 {len(names)} 条 cookie，含登录凭证 "
          f"({'sessionid' if 'sessionid' in names else 'sessionid_ss'})")

    # 2. 表格
    step(2, "从表格里取作品 id ...")
    aweme_id, stored = find_row(xlsx)
    print(f"      作品 id：{aweme_id}")
    if stored:
        exp = douyin_api.url_expiry(str(stored))
        if exp:
            hrs = (time.time() - exp) / 3600
            state = f"已过期 {hrs:.1f} 小时" if hrs > 0 else f"还有 {-hrs:.1f} 小时到期"
            print(f"      表格里存的地址：{state}")

    # 3. 现取
    step(3, "调抖音接口现取新地址 ...")
    try:
        api = douyin_api.DouyinAPI(header)
    except douyin_api.DouyinApiError as e:
        die(f"{e}")
    try:
        fresh = api.fresh_play_url(aweme_id)
    except douyin_api.DouyinApiError as e:
        die(f"现取失败：{e}\n"
            "   → 最常见原因是 cookies 过期。重新登录抖音、重新导出 cookies.txt 再试。")
    except Exception as e:  # noqa: BLE001
        die(f"现取时出错：{type(e).__name__}: {e}")

    exp = douyin_api.url_expiry(fresh)
    if exp:
        print(f"      拿到新地址，{(exp - time.time()) / 3600:.1f} 小时后到期")
    else:
        print("      拿到新地址")
    print(f"      {fresh[:95]}...")

    # 4. 真下载
    step(4, "试着下载这个新地址 ...")
    code, info = probe(fresh)
    print(f"      HTTP {code}  {info}")

    print()
    if code == 200:
        print("✅ 全链路通了。回到界面里，「Cookies 文件」选这份 cookies.txt，就能正常跑整批。")
    elif code == 403:
        print("❌ 新地址也被拒。说明这个 cookie 权限不够（比如没真正登录），"
              "重新登录抖音后重新导出 cookies.txt。")
    elif code is None:
        print(f"❌ 连不上：{info}")
    else:
        print(f"❌ 返回 {code}，把这个数字发我。")


if __name__ == "__main__":
    main()
