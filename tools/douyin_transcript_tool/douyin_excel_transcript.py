#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 Excel 表格批量读取抖音作品链接，转写后写回同一份表格

专门配合"抖音主页/热点宝作品采集"类工具导出的 Excel（列名里带"作品网址"、
"是否已经语音转写文案"这种字段的表格），不需要再把链接手动复制到 links.txt 里。

用法示例：
    # 本地模式
    python3 douyin_excel_transcript.py --excel-file 作品列表.xlsx --mode local

    # SiliconFlow（api 模式 + 自定义接口地址）
    python3 douyin_excel_transcript.py --excel-file 作品列表.xlsx \
        --mode api --api-key sk-xxxx --api-base-url https://api.siliconflow.cn/v1 \
        --api-model FunAudioLLM/SenseVoiceSmall

依赖：和 douyin_batch_transcript.py 用同一套依赖（yt-dlp / openai-whisper 或 openai /
requests / dashscope 等，具体看你用哪个转写模式），另外还需要 openpyxl。
douyin_excel_transcript.py 需要和 douyin_batch_transcript.py 放在同一个文件夹里，
是直接复用后者的下载/转写逻辑，不是另一套实现。

断点续传：脚本会用"是否已经语音转写文案"这一列判断某一行是否已经处理过，已经是
True 的行会自动跳过，重复运行不会重复转写。处理过程中每处理完一行就会保存一次
文件，即使中途报错/中断，已经转写完的部分也不会丢。

安全提醒：这个脚本会直接修改你传进来的 Excel 文件本身（原地追加/更新列），
建议第一次用之前自己先备份一份。
"""

import argparse
import logging
import os
import random
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import douyin_batch_transcript as core
except ImportError as exc:
    raise RuntimeError(
        "找不到 douyin_batch_transcript.py，请确认它和本文件在同一个文件夹里。"
    ) from exc

try:
    import openpyxl
except ImportError:
    openpyxl = None

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("douyin_excel_transcript")

_DONE_VALUES = {"true", "是", "1", "yes"}


def _is_marked_done(value) -> bool:
    if value is True:
        return True
    if value is None:
        return False
    return str(value).strip().lower() in _DONE_VALUES


def find_column_index(header_row, name):
    for idx, cell in enumerate(header_row, start=1):
        if cell == name:
            return idx
    return None


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="从 Excel 批量读取抖音链接并把逐字稿写回表格",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--excel-file", required=True, help="Excel 文件路径（会直接在原文件里追加/更新列，建议先备份）")
    parser.add_argument("--sheet", default="", help="工作表名称，不填默认用第一个工作表")
    parser.add_argument("--link-column", default="作品网址", help="视频链接所在的列名")
    parser.add_argument("--status-column", default="是否已经语音转写文案", help="标记是否已转写的列名，用于断点续传，不存在会自动新增")
    parser.add_argument("--transcript-column", default="逐字稿", help="写入转写文字的列名，不存在会自动新增")
    parser.add_argument("--mode", choices=["local", "api", "bailian"], default="local", help="转写模式，含义同 douyin_batch_transcript.py")
    parser.add_argument("--model-size", default="small", choices=["tiny", "base", "small", "medium", "large"])
    parser.add_argument("--api-key", default=os.environ.get("OPENAI_API_KEY", ""))
    parser.add_argument("--api-base-url", default="")
    parser.add_argument("--api-model", default="whisper-1")
    parser.add_argument("--dashscope-api-key", default=os.environ.get("DASHSCOPE_API_KEY", ""))
    parser.add_argument("--cookies-from-browser", default="", help="从指定浏览器读取 Cookies，例如 chrome/safari/firefox/edge")
    parser.add_argument("--cookies-file", default="", help="Netscape 格式的 cookies.txt 文件路径")
    parser.add_argument("--tmp-dir", default="", help="下载音频的临时目录，默认在 Excel 文件同目录下建一个 _audio_tmp")
    parser.add_argument(
        "--keep-video", action="store_true",
        help="同时保留下载到的原始视频文件（保存到 Excel 同目录下的 <视频ID>.mp4）",
    )
    parser.add_argument("--min-delay", type=float, default=1.0, help="每行开始前的最小随机延迟（秒）")
    parser.add_argument("--max-delay", type=float, default=3.0, help="每行开始前的最大随机延迟（秒）")
    parser.add_argument("--retries", type=int, default=3, help="每步失败后的最大重试次数")
    parser.add_argument("--retry-base-delay", type=float, default=2.0, help="重试的基础等待秒数（指数递增）")
    return parser


def main():
    args = build_arg_parser().parse_args()

    if openpyxl is None:
        logger.error("未安装 openpyxl，请先运行: pip install openpyxl")
        sys.exit(1)

    if not os.path.isfile(args.excel_file):
        logger.error("找不到 Excel 文件：%s，请确认路径是否正确。", args.excel_file)
        sys.exit(1)

    if args.mode == "api" and not args.api_key:
        logger.error("api 模式下必须提供 --api-key，或设置环境变量 OPENAI_API_KEY。")
        sys.exit(1)
    if args.mode == "bailian" and not args.dashscope_api_key:
        logger.error("bailian 模式下必须提供 --dashscope-api-key，或设置环境变量 DASHSCOPE_API_KEY。")
        sys.exit(1)

    tmp_dir = args.tmp_dir or os.path.join(os.path.dirname(os.path.abspath(args.excel_file)) or ".", "_audio_tmp")
    try:
        os.makedirs(tmp_dir, exist_ok=True)
    except OSError as exc:
        logger.error("无法创建临时音频目录 %s：%s", tmp_dir, exc)
        sys.exit(1)

    try:
        wb = openpyxl.load_workbook(args.excel_file)
    except Exception as exc:  # noqa: BLE001 - openpyxl 对各种损坏/加密文件抛出的异常类型不统一
        logger.error("打开 Excel 文件失败：%s", exc)
        sys.exit(1)

    if args.sheet:
        if args.sheet not in wb.sheetnames:
            logger.error("找不到工作表 '%s'，实际的工作表有：%s", args.sheet, wb.sheetnames)
            sys.exit(1)
        ws = wb[args.sheet]
    else:
        ws = wb.worksheets[0]

    header = [cell.value for cell in ws[1]]
    link_col = find_column_index(header, args.link_column)
    if link_col is None:
        logger.error(
            "表格 '%s' 里找不到链接列 '%s'。表格实际的列名有：%s",
            ws.title, args.link_column, header,
        )
        sys.exit(1)

    status_col = find_column_index(header, args.status_column)
    if status_col is None:
        status_col = len(header) + 1
        ws.cell(row=1, column=status_col, value=args.status_column)
        header.append(args.status_column)
        logger.info("表格里没有 '%s' 列，已自动新增。", args.status_column)

    transcript_col = find_column_index(header, args.transcript_column)
    if transcript_col is None:
        transcript_col = len(header) + 1
        ws.cell(row=1, column=transcript_col, value=args.transcript_column)
        header.append(args.transcript_column)
        logger.info("表格里没有 '%s' 列，已自动新增。", args.transcript_column)

    total_rows = ws.max_row - 1
    pending_rows = []
    skipped = 0
    for row_idx in range(2, ws.max_row + 1):
        link = ws.cell(row=row_idx, column=link_col).value
        if not link or not str(link).strip():
            continue
        if _is_marked_done(ws.cell(row=row_idx, column=status_col).value):
            skipped += 1
            continue
        pending_rows.append(row_idx)

    if skipped:
        logger.info("检测到断点续传记录，跳过已标记完成的 %d 行。", skipped)

    if not pending_rows:
        logger.info("没有待处理的行（共 %d 行数据）。", total_rows)
        return

    logger.info("共 %d 行数据，待处理 %d 行，模式=%s", total_rows, len(pending_rows), args.mode)

    success = 0
    failed = 0
    for done_count, row_idx in enumerate(pending_rows, start=1):
        link = str(ws.cell(row=row_idx, column=link_col).value).strip()
        logger.info("[%d/%d] 第 %d 行，开始处理：%s", done_count, len(pending_rows), row_idx, link)

        time.sleep(random.uniform(args.min_delay, args.max_delay))

        audio_path = None
        save_dir = os.path.dirname(os.path.abspath(args.excel_file)) or "."
        try:
            audio_path, video_id, _title, thumbnail_url, video_path = core.retry_with_backoff(
                lambda: core.download_audio(
                    link, tmp_dir,
                    cookies_from_browser=args.cookies_from_browser,
                    cookies_file=args.cookies_file,
                    keep_video_dir=save_dir if args.keep_video else None,
                ),
                retries=args.retries, base_delay=args.retry_base_delay,
                what="下载音频", link=link,
            )
            if video_path:
                logger.info("[%d/%d] 第 %d 行：视频文件已保留到 %s", done_count, len(pending_rows), row_idx, video_path)
            text, segments = core.retry_with_backoff(
                lambda: core.transcribe_dispatch(audio_path, args.mode, args),
                retries=args.retries, base_delay=args.retry_base_delay,
                what=f"{args.mode}转写", link=link,
            )
            ws.cell(row=row_idx, column=transcript_col, value=core.clean_for_excel(text))
            ws.cell(row=row_idx, column=status_col, value="True")
            success += 1
            logger.info("[%d/%d] 第 %d 行成功。", done_count, len(pending_rows), row_idx)

            core.download_cover_image(thumbnail_url, save_dir, video_id)
        except Exception as exc:  # noqa: BLE001
            ws.cell(row=row_idx, column=transcript_col, value=core.clean_for_excel(f"转写失败：{exc}"))
            failed += 1
            logger.error("[%d/%d] 第 %d 行失败：%s", done_count, len(pending_rows), row_idx, exc)
        finally:
            if audio_path and os.path.isfile(audio_path):
                try:
                    os.remove(audio_path)
                except OSError:
                    pass

        try:
            wb.save(args.excel_file)
        except OSError as exc:
            logger.error("保存 Excel 文件失败（可能文件正被 Excel/WPS 打开着）：%s", exc)
            sys.exit(1)

    logger.info("全部处理完成：成功 %d 条，失败 %d 条。", success, failed)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        logger.warning("用户中断，已处理的部分结果已经保存在 Excel 文件里，不会丢失。")
        sys.exit(130)
